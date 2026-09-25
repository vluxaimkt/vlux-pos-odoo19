"""The catalog template, generated for the caller's company.

A static file could not list the store's own taxes and categories, and a
wrong tax name is the most common import error; the template carries them
in two reference sheets.
"""
import io

from odoo import http
from odoo.exceptions import AccessError
from odoo.http import request

from odoo.addons.vlux_pos_catalog.models.catalog_import import IMPORT_GROUPS
from odoo.addons.vlux_pos_catalog.models.catalog_import_parser import TEMPLATE_HEADERS

EXAMPLES = (
    ("7501055300075", "Refresco cola 600 ml", 18.5, 11.2, "REF-600", "Bebidas", "", "default", 24, "si", "si"),
    ("7501000112346", "Galletas surtidas 200 g", 32, 21, "", "Abarrotes / Galletas", "", "", 10, "si", "si"),
)
NOTES = (
    ("codigo_barras", "Identifica el producto. Si ya existe, la fila lo actualiza; si no, lo crea."),
    ("nombre", "Obligatorio para productos nuevos."),
    ("precio_venta", "Obligatorio para productos nuevos. Acepta 18.50, 18,50 o $1,234.50."),
    ("costo", "Opcional."),
    ("referencia", "Opcional. Identifica el producto cuando no tiene código de barras."),
    ("categoria_pos", "Botón del punto de venta. Usa / para subcategorías: Abarrotes / Galletas."),
    ("categoria", "Categoría interna (contable). Opcional."),
    ("impuestos", "Nombre exacto de la hoja Impuestos. Vacío = impuesto por defecto de la empresa. "
                  "Varios: IEPS 8% + IVA 16%. Sin impuesto: ninguno."),
    ("existencia", "Cantidad en inventario. Reemplaza la existencia actual (no suma)."),
    ("inventariable", "si / no. Por defecto sí cuando hay existencia."),
    ("disponible_pos", "si / no. Por defecto sí."),
    ("", "Una celda vacía deja ese dato como está al actualizar."),
)


class VluxCatalogImportController(http.Controller):

    @http.route("/vlux_pos_catalog/import/template.xlsx", type="http", auth="user", methods=["GET"])
    def template(self, **kwargs):
        user = request.env.user
        if not any(user.has_group(group) for group in IMPORT_GROUPS):
            raise AccessError("Sin permiso para importar el catálogo.")
        import openpyxl
        from openpyxl.styles import Font, PatternFill

        company = request.env.company
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Productos"
        sheet.append(list(TEMPLATE_HEADERS))
        default_tax = company.account_sale_tax_id[:1].name or ""
        for example in EXAMPLES:
            row = list(example)
            row[7] = default_tax if row[7] else ""  # a tax name that exists in this company
            sheet.append(row)
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="6B4E9B")
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
        for column, width in zip("ABCDEFGHIJK", (17, 32, 13, 10, 13, 24, 18, 18, 11, 13, 14)):
            sheet.column_dimensions[column].width = width
        # Barcodes are text: keep Excel from turning them into 7.50106E+12.
        for row in sheet.iter_rows(min_row=2, max_row=5000, min_col=1, max_col=1):
            for cell in row:
                cell.number_format = "@"
        sheet.freeze_panes = "A2"

        taxes = request.env["account.tax"].search(
            [("type_tax_use", "=", "sale"), ("company_id", "=", company.id)], order="sequence, name"
        )
        tax_sheet = book.create_sheet("Impuestos")
        tax_sheet.append(["nombre (cópialo tal cual)", "porcentaje", "por defecto"])
        for tax in taxes:
            tax_sheet.append([tax.name, tax.amount, "sí" if tax in company.account_sale_tax_id else ""])
        tax_sheet.column_dimensions["A"].width = 34

        category_sheet = book.create_sheet("Categorias")
        category_sheet.append(["categoria_pos existentes", "categoria existentes"])
        pos_categories = request.env["pos.category"].search([], order="name")
        categories = request.env["product.category"].search([], order="complete_name")
        for index in range(max(len(pos_categories), len(categories))):
            category_sheet.append([
                pos_categories[index].display_name if index < len(pos_categories) else "",
                categories[index].complete_name if index < len(categories) else "",
            ])
        category_sheet.column_dimensions["A"].width = 32
        category_sheet.column_dimensions["B"].width = 40

        help_sheet = book.create_sheet("Instrucciones")
        help_sheet.append(["columna", "cómo llenarla"])
        for note in NOTES:
            help_sheet.append(list(note))
        help_sheet.column_dimensions["A"].width = 16
        help_sheet.column_dimensions["B"].width = 110
        for header_row in (tax_sheet[1], category_sheet[1], help_sheet[1]):
            for cell in header_row:
                cell.font = Font(bold=True)

        stream = io.BytesIO()
        book.save(stream)
        return request.make_response(stream.getvalue(), headers=[
            ("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            ("Content-Disposition", 'attachment; filename="plantilla_catalogo_vlux.xlsx"'),
            ("Cache-Control", "no-store"),
        ])
