"""Reading a store's catalog spreadsheet into clean rows.

Pure Python on purpose (no ORM): the same rules are unit-tested without a
database, and the model only resolves what needs the database (taxes,
categories, existing products).

Spreadsheets arrive from Excel, from a supplier's export or from someone's
phone, so the reader is forgiving about format and strict about meaning:
headers match in Spanish or English, with or without accents; numbers may
carry ``$`` or thousands separators; a decimal comma is accepted; anything
it cannot read is reported per row, never guessed.
"""
import csv
import io
import math
import re
import unicodedata

MAX_ROWS = 50_000
MAX_BYTES = 20 * 1024 * 1024
BARCODE_RE = re.compile(r"^[A-Za-z0-9\-_.]{3,64}$")
NAME_MAX_LENGTH = 256
REFERENCE_MAX_LENGTH = 64
PRICE_MAX = 10_000_000.0
QTY_MAX = 1_000_000.0

# canonical column -> accepted header spellings (already normalised)
COLUMNS = {
    "barcode": ("codigo_barras", "codigo_de_barras", "codigo", "barcode", "ean", "upc", "gtin"),
    "name": ("nombre", "producto", "descripcion", "name"),
    "list_price": ("precio_venta", "precio", "precio_de_venta", "list_price", "price"),
    "standard_price": ("costo", "cost", "precio_compra", "standard_price"),
    "default_code": ("referencia", "referencia_interna", "sku", "clave", "default_code"),
    "pos_category": ("categoria_pos", "categoria_punto_de_venta", "pos_category"),
    "category": ("categoria", "categoria_producto", "category"),
    "taxes": ("impuestos", "impuesto", "iva", "taxes", "tax"),
    "quantity": ("existencia", "existencias", "stock", "stock_inicial", "cantidad", "inventario", "quantity"),
    "is_storable": ("inventariable", "controlar_inventario", "is_storable"),
    "available_in_pos": ("disponible_pos", "disponible_en_pos", "vender_en_pos", "available_in_pos"),
}
TEMPLATE_HEADERS = (
    "codigo_barras", "nombre", "precio_venta", "costo", "referencia",
    "categoria_pos", "categoria", "impuestos", "existencia", "inventariable", "disponible_pos",
)
TRUE_WORDS = {"si", "s", "yes", "y", "1", "true", "verdadero", "x"}
FALSE_WORDS = {"no", "n", "0", "false", "falso"}


class ParseError(Exception):
    """The file itself cannot be read (as opposed to one bad row)."""


def normalize(text):
    """Lowercase, no accents, words joined by ``_``: how headers are compared."""
    text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


def _column_index(headers):
    """Map canonical columns to their position; unknown headers are ignored."""
    aliases = {alias: column for column, spellings in COLUMNS.items() for alias in spellings}
    index, seen = {}, set()
    for position, header in enumerate(headers):
        column = aliases.get(normalize(header))
        if column and column not in seen:
            index[column] = position
            seen.add(column)
    if "barcode" not in index and "default_code" not in index:
        raise ParseError("Falta la columna codigo_barras (o referencia): es la que identifica cada producto.")
    if len(index) == 1:
        raise ParseError("El archivo sólo trae la columna que identifica al producto; no hay nada que importar.")
    # No other column is mandatory: a file with codigo_barras + precio_venta
    # is a price update. A *new* product without nombre or precio_venta is
    # reported on its own row by the model.
    return index


def read_table(data, filename):
    """Return ``(headers, rows)`` from an XLSX or CSV payload."""
    if not data:
        raise ParseError("El archivo está vacío.")
    if len(data) > MAX_BYTES:
        raise ParseError("El archivo supera %d MB." % (MAX_BYTES // (1024 * 1024)))
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
        return _read_xlsx(data)
    if name.endswith((".csv", ".txt")) or not name:
        return _read_csv(data)
    raise ParseError("Formato no soportado: usa .xlsx o .csv.")


def _read_xlsx(data):
    try:
        import openpyxl
    except ImportError:  # pragma: no cover - Odoo ships openpyxl
        raise ParseError("El servidor no puede leer archivos .xlsx; usa .csv.")
    try:
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - any failure means "not a readable workbook"
        raise ParseError("No se pudo abrir el archivo de Excel.")
    sheet = book["Productos"] if "Productos" in book.sheetnames else book.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows, None)
    if not headers:
        raise ParseError("La hoja no tiene encabezados.")
    body = []
    for row in rows:
        body.append(list(row))
        if len(body) > MAX_ROWS:
            raise ParseError("El archivo tiene más de %d filas." % MAX_ROWS)
    book.close()
    return [str(h or "") for h in headers], body


def _read_csv(data):
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    first_line = text.split("\n", 1)[0]
    delimiter = max((";", ",", "\t", "|"), key=first_line.count)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    headers = next(reader, None)
    if not headers:
        raise ParseError("El archivo no tiene encabezados.")
    body = []
    for row in reader:
        body.append(row)
        if len(body) > MAX_ROWS:
            raise ParseError("El archivo tiene más de %d filas." % MAX_ROWS)
    return headers, body


def _blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def parse_number(value):
    """``12``, ``12.5``, ``"$1,234.50"``, ``"1.234,50"``, ``"12,5"`` -> float.

    Raises ValueError when the text is not a number. With both separators the
    last one is the decimal mark; with only a comma, it is the decimal mark
    unless it groups exactly three digits more than once (``1,234,567``).
    """
    if isinstance(value, bool):
        raise ValueError(value)
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = str(value).strip().replace("$", "").replace(" ", "").replace(" ", "")
        if text.startswith("(") and text.endswith(")"):
            text = "-" + text[1:-1]
        if "," in text and "." in text:
            if text.rfind(",") > text.rfind("."):
                text = text.replace(".", "").replace(",", ".")
            else:
                text = text.replace(",", "")
        elif "," in text:
            if re.fullmatch(r"-?\d{1,3}(,\d{3}){2,}", text):
                text = text.replace(",", "")
            else:
                text = text.replace(",", ".")
        number = float(text)
    if not math.isfinite(number):
        raise ValueError(value)
    return number


def parse_bool(value):
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    word = normalize(value)
    if word in TRUE_WORDS:
        return True
    if word in FALSE_WORDS:
        return False
    raise ValueError(value)


def _text(value):
    if _blank(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        # Excel turns 7501234567890 into 7501234567890.0
        return str(int(value))
    return str(value).strip()


def parse_rows(headers, rows):
    """Clean every row. Returns a list of ``{"row": n, "values": {...}, "errors": [...]}``.

    ``n`` is the spreadsheet row number (header = 1). ``values`` only holds the
    columns that were filled in: an empty cell means "leave as is" on update.
    Fully empty rows are skipped. Resolution of taxes and categories happens
    later, in the model.
    """
    index = _column_index(headers)
    parsed = []
    seen_keys = {}
    for offset, raw in enumerate(rows):
        number = offset + 2
        cell = {column: (raw[position] if position < len(raw) else None) for column, position in index.items()}
        if all(_blank(value) for value in cell.values()):
            continue
        values, errors = {}, []

        barcode = _text(cell.get("barcode"))
        if barcode:
            if BARCODE_RE.match(barcode):
                values["barcode"] = barcode
            else:
                errors.append(("codigo_barras", "Código de barras inválido: 3 a 64 letras, números, - _ ."))
        reference = _text(cell.get("default_code"))
        if reference:
            if len(reference) > REFERENCE_MAX_LENGTH:
                errors.append(("referencia", "La referencia supera %d caracteres." % REFERENCE_MAX_LENGTH))
            else:
                values["default_code"] = reference
        if not barcode and not reference:
            errors.append(("codigo_barras", "Cada producto necesita código de barras o referencia para identificarlo."))

        name = _text(cell.get("name"))
        if name:
            if len(name) > NAME_MAX_LENGTH:
                errors.append(("nombre", "El nombre supera %d caracteres." % NAME_MAX_LENGTH))
            else:
                values["name"] = name

        for column, label, maximum in (
            ("list_price", "precio_venta", PRICE_MAX),
            ("standard_price", "costo", PRICE_MAX),
            ("quantity", "existencia", QTY_MAX),
        ):
            value = cell.get(column)
            if _blank(value):
                continue
            try:
                amount = parse_number(value)
            except (TypeError, ValueError):
                errors.append((label, "«%s» no es un número." % _text(value)))
                continue
            if amount < 0 or amount > maximum:
                errors.append((label, "%s fuera de rango (0 a %s)." % (label, f"{maximum:,.0f}")))
                continue
            values[column] = amount

        for column, label in (("is_storable", "inventariable"), ("available_in_pos", "disponible_pos")):
            value = cell.get(column)
            if _blank(value):
                continue
            try:
                values[column] = parse_bool(value)
            except ValueError:
                errors.append((label, "«%s» no es sí/no." % _text(value)))

        for column in ("pos_category", "category", "taxes"):
            text = _text(cell.get(column))
            if text:
                values[column] = text

        if values.get("is_storable") is False and "quantity" in values:
            errors.append(("existencia", "Un producto no inventariable no lleva existencia."))

        key = ("barcode", barcode) if barcode else ("default_code", reference)
        if key[1]:
            first = seen_keys.setdefault(key, number)
            if first != number:
                label = "código de barras" if key[0] == "barcode" else "referencia"
                errors.append(("codigo_barras" if key[0] == "barcode" else "referencia",
                               "El %s %s ya aparece en la fila %d." % (label, key[1], first)))
        parsed.append({"row": number, "values": values, "errors": errors})
    return parsed
