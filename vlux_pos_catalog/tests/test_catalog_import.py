import base64
import io
from unittest.mock import patch

from odoo.exceptions import AccessError, UserError
from odoo.tests.common import HttpCase, tagged

from ..models import catalog_import
from ..models import catalog_import_parser as parser
from .common import VluxCatalogCase

HEADERS = list(parser.TEMPLATE_HEADERS)


def xlsx(rows, headers=HEADERS):
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Productos"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def csv_bytes(text, encoding="utf-8"):
    return text.encode(encoding)


@tagged("post_install", "-at_install")
class TestCatalogImportParser(VluxCatalogCase):
    """The spreadsheet reader: forgiving about format, strict about meaning."""

    def test_numbers_in_the_formats_stores_actually_use(self):
        for text, expected in (("18.5", 18.5), ("18,50", 18.5), ("$1,234.50", 1234.5), ("1.234,50", 1234.5),
                               ("1,234,567", 1234567), (" 12 ", 12), (7, 7.0), (3.25, 3.25)):
            self.assertAlmostEqual(parser.parse_number(text), expected, msg=text)
        for bad in ("abc", "1,2,3.4.5", "", "nan", "inf", True):
            with self.assertRaises(ValueError, msg=bad):
                parser.parse_number(bad)

    def test_headers_match_in_spanish_or_english_with_or_without_accents(self):
        headers = ["Código de Barras", "DESCRIPCIÓN", "Precio", "IVA", "Stock"]
        parsed = parser.parse_rows(headers, [["7501", "Agua", "10", "16", "5"]])
        values = parsed[0]["values"]
        self.assertEqual((values["barcode"], values["name"], values["list_price"], values["taxes"], values["quantity"]),
                         ("7501", "Agua", 10.0, "16", 5.0))
        with self.assertRaises(parser.ParseError):
            parser.parse_rows(["nombre", "precio"], [["Agua", "10"]])
        with self.assertRaises(parser.ParseError):
            parser.parse_rows(["codigo_barras", "columna_desconocida"], [["7501", "10"]])
        # codigo_barras + precio is a valid price-update file
        self.assertEqual(parser.parse_rows(["codigo_barras", "precio"], [["7501", "10"]])[0]["values"],
                         {"barcode": "7501", "list_price": 10.0})

    def test_row_errors_are_specific_and_blank_rows_are_skipped(self):
        rows = [
            ["7501000000001", "Bueno", "10", "", "", "", "", "", "", "", ""],
            ["", "", "", "", "", "", "", "", "", "", ""],
            ["7501 000", "Espacio en el código", "10", "", "", "", "", "", "", "", ""],
            ["7501000000003", "Precio malo", "diez", "", "", "", "", "", "", "", ""],
            ["7501000000001", "Duplicado", "10", "", "", "", "", "", "", "", ""],
            ["", "Sin código", "10", "", "", "", "", "", "", "", ""],
            ["7501000000004", "Negativo", "-1", "", "", "", "", "", "", "", ""],
            ["7501000000005", "No inventariable con stock", "1", "", "", "", "", "", "3", "no", ""],
            ["7501000000006", "Booleano raro", "1", "", "", "", "", "", "", "tal vez", ""],
        ]
        parsed = parser.parse_rows(HEADERS, rows)
        by_row = {item["row"]: item for item in parsed}
        self.assertNotIn(3, by_row, "a fully empty row is skipped")
        self.assertEqual(by_row[2]["errors"], [])
        self.assertEqual(by_row[4]["errors"][0][0], "codigo_barras")
        self.assertIn("no es un número", by_row[5]["errors"][0][1])
        self.assertIn("ya aparece en la fila 2", by_row[6]["errors"][0][1])
        self.assertIn("código de barras o referencia", by_row[7]["errors"][0][1])
        self.assertIn("fuera de rango", by_row[8]["errors"][0][1])
        self.assertIn("no inventariable", by_row[9]["errors"][0][1])
        self.assertIn("sí/no", by_row[10]["errors"][0][1])

    def test_csv_with_semicolons_and_windows_encoding_and_excel_float_barcodes(self):
        text = "codigo_barras;nombre;precio_venta\n7501000000011;Café molido;45,90\n"
        headers, rows = parser.read_table(csv_bytes(text, "cp1252"), "catalogo.csv")
        values = parser.parse_rows(headers, rows)[0]["values"]
        self.assertEqual((values["name"], values["list_price"]), ("Café molido", 45.9))

        headers, rows = parser.read_table(xlsx([[7501000000012.0, "Leche", 25]]), "catalogo.xlsx")
        self.assertEqual(parser.parse_rows(headers, rows)[0]["values"]["barcode"], "7501000000012")

        with self.assertRaises(parser.ParseError):
            parser.read_table(b"", "vacio.csv")
        with self.assertRaises(parser.ParseError):
            parser.read_table(b"%PDF-1.4", "catalogo.pdf")


@tagged("post_install", "-at_install")
class TestCatalogImport(VluxCatalogCase):
    """Validate writes nothing; import creates or updates, idempotently."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.tax16 = cls.env["account.tax"].sudo().create({
            "name": "VLUX IVA 16%", "amount": 16, "type_tax_use": "sale", "company_id": cls.company_a.id,
        })
        cls.tax8 = cls.env["account.tax"].sudo().create({
            "name": "VLUX IEPS 8%", "amount": 8, "type_tax_use": "sale", "company_id": cls.company_a.id,
        })
        cls.location = cls.env["stock.warehouse"].search([("company_id", "=", cls.company_a.id)], limit=1).lot_stock_id

    def new_import(self, data, filename="catalogo.xlsx", user=None, **values):
        user = user or self.owner
        return self.env["vlux.catalog.import"].with_user(user).with_company(user.company_id).create({
            "name": filename, "file": base64.b64encode(data), "location_id": self.location.id, **values,
        })

    def run_import(self, record):
        record.action_validate()
        record.action_import()
        return record.sudo()._run_to_completion()

    def product(self, barcode):
        return self.env["product.product"].with_context(active_test=False).search([("barcode", "=", barcode)])

    def test_validate_reports_counts_and_errors_without_touching_the_catalog(self):
        existing = self.env["product.template"].create({
            "name": "Ya existe", "barcode": "7502000000001", "list_price": 5, "company_id": self.company_a.id,
        })
        foreign = self.env["product.template"].sudo().create({
            "name": "De otra empresa", "barcode": "7502000000099", "list_price": 5, "company_id": self.company_b.id,
        })
        count_before = self.env["product.template"].search_count([])
        data = xlsx([
            ["7502000000001", "", "7.5", "", "", "", "", "", "", "", ""],           # update: price only
            ["7502000000002", "Nuevo", "12", "", "", "Bebidas", "", "VLUX IVA 16%", "", "", ""],
            ["7502000000003", "", "12", "", "", "", "", "", "", "", ""],             # new without name
            ["7502000000004", "Impuesto raro", "12", "", "", "", "", "IVA 99%", "", "", ""],
            ["7502000000099", "Ajeno", "12", "", "", "", "", "", "", "", ""],
        ])
        record = self.new_import(data)
        record.action_validate()

        self.assertEqual(record.state, "validated")
        self.assertEqual((record.rows_total, record.rows_valid, record.rows_to_create, record.rows_to_update), (5, 2, 1, 1))
        messages = {error.row: error.message for error in record.error_ids}
        self.assertIn("falta nombre", messages[4])
        self.assertIn("No existe el impuesto", messages[5])
        self.assertIn("otra empresa", messages[6])
        self.assertEqual(self.env["product.template"].search_count([]), count_before, "validation writes nothing")
        self.assertEqual(existing.list_price, 5)
        self.assertEqual(foreign.sudo().name, "De otra empresa")

    def test_import_creates_with_taxes_categories_and_stock(self):
        data = xlsx([
            ["7502000000101", "Refresco 600", "18.5", "11", "REF-1", "Bebidas / Refrescos", "", "VLUX IEPS 8% + VLUX IVA 16%", "24", "", ""],
            ["7502000000102", "Bolsa", "2", "", "", "", "", "ninguno", "", "no", "si"],
            ["7502000000103", "Por monto", "5", "", "", "", "", "16", "", "", "no"],
        ])
        record = self.run_import(self.new_import(data))

        self.assertEqual(record.state, "done")
        self.assertEqual((record.created_count, record.updated_count, record.stock_count, record.error_count), (3, 0, 1, 0))
        soda = self.product("7502000000101")
        self.assertEqual((soda.name, soda.lst_price, soda.standard_price, soda.default_code), ("Refresco 600", 18.5, 11, "REF-1"))
        self.assertEqual(set(soda.taxes_id.ids), {self.tax8.id, self.tax16.id})
        self.assertEqual(soda.pos_categ_ids.name, "Refrescos")
        self.assertEqual(soda.pos_categ_ids.parent_id.name, "Bebidas")
        self.assertTrue(soda.is_storable and soda.available_in_pos)
        self.assertEqual(soda.company_id, self.company_a)
        self.assertEqual(soda.with_context(location=self.location.id).qty_available, 24)
        bag = self.product("7502000000102")
        self.assertFalse(bag.taxes_id)
        self.assertFalse(bag.is_storable)
        by_amount = self.product("7502000000103")
        self.assertEqual(by_amount.taxes_id, self.tax16)
        self.assertFalse(by_amount.available_in_pos)

    def test_import_is_an_idempotent_upsert(self):
        data = xlsx([["7502000000201", "Agua 1L", "10", "", "", "", "", "", "5", "", ""]])
        self.run_import(self.new_import(data))
        product = self.product("7502000000201")
        self.assertEqual(product.qty_available, 5)

        # Same file again: no duplicate, nothing rewritten, stock left at 5 (not added).
        stamp = product.vlux_sync_date
        moves = self.env["stock.move"].search_count([("product_id", "=", product.id)])
        second = self.run_import(self.new_import(data))
        self.assertEqual((second.created_count, second.updated_count, second.unchanged_count, second.stock_count),
                         (0, 0, 1, 0))
        self.assertEqual(len(self.product("7502000000201")), 1)
        self.assertEqual(product.qty_available, 5)
        self.assertEqual(product.vlux_sync_date, stamp, "an unchanged product does not bounce every register's sync")
        self.assertEqual(self.env["stock.move"].search_count([("product_id", "=", product.id)]), moves)

        # A price-only file (two columns) leaves name, taxes and stock alone.
        taxes = product.taxes_id
        third = self.run_import(self.new_import(xlsx([["7502000000201", "11.5"]], headers=["codigo_barras", "precio_venta"])))
        self.assertEqual(third.updated_count, 1)
        self.assertEqual((product.name, product.lst_price, product.taxes_id, product.qty_available), ("Agua 1L", 11.5, taxes, 5))

    def test_archived_products_come_back_and_reference_is_a_key_without_barcode(self):
        archived = self.env["product.template"].create({
            "name": "Archivado", "barcode": "7502000000301", "list_price": 3, "available_in_pos": False,
            "company_id": self.company_a.id,
        })
        archived.action_archive()
        by_reference = self.env["product.template"].create({
            "name": "Sin código", "default_code": "GRANEL-1", "list_price": 1, "company_id": self.company_a.id,
        })
        data = xlsx([
            ["7502000000301", "", "4", "", "", "", "", "", "", "", ""],
            ["", "Frijol a granel", "30", "", "GRANEL-1", "", "", "", "", "", ""],
        ])
        record = self.run_import(self.new_import(data))
        self.assertEqual((record.created_count, record.updated_count), (0, 2))
        self.assertTrue(archived.active)
        self.assertEqual(archived.list_price, 4)
        self.assertEqual((by_reference.name, by_reference.list_price), ("Frijol a granel", 30))

    def test_missing_categories_can_be_refused(self):
        data = xlsx([["7502000000401", "Algo", "1", "", "", "Categoría inventada", "", "", "", "", ""]])
        record = self.new_import(data, create_missing_categories=False)
        record.action_validate()
        self.assertEqual(record.rows_valid, 0)
        self.assertIn("no existe", record.error_ids.message)
        with self.assertRaises(UserError):
            record.action_import()

    def test_a_failing_row_does_not_sink_its_batch(self):
        data = xlsx([[f"750200000050{i}", f"Producto {i}", "1", "", "", "", "", "", "", "", ""] for i in range(5)])
        original = catalog_import._Importer._template_values

        def flaky(importer, values, creating):
            if values.get("barcode") == "7502000000503":
                raise UserError("fila venenosa")
            return original(importer, values, creating)

        record = self.new_import(data)
        record.action_validate()
        record.action_import()
        with patch.object(catalog_import._Importer, "_template_values", flaky):
            record.sudo()._run_to_completion()
        self.assertEqual(record.state, "done")
        self.assertEqual(record.created_count, 4)
        self.assertEqual(record.error_ids.mapped("row"), [5])
        self.assertIn("fila venenosa", record.error_ids.message)
        self.assertFalse(self.product("7502000000503"))

    def test_batches_resume_where_they_stopped(self):
        data = xlsx([[f"75020000006{i:02d}", f"Lote {i}", "1", "", "", "", "", "", "", "", ""] for i in range(7)])
        record = self.new_import(data)
        record.action_validate()
        record.action_import()
        with patch.object(catalog_import, "BATCH_SIZE", 3):
            self.assertFalse(record.sudo()._run_batch())
            self.assertEqual((record.state, record.rows_done), ("running", 3))
            # A replayed batch (crash before the progress was saved) creates nothing twice.
            record.sudo().rows_done = 0
            record.sudo()._run_to_completion()
        self.assertEqual(record.state, "done")
        for i in range(7):
            self.assertEqual(len(self.product(f"75020000006{i:02d}")), 1)

    def test_only_catalog_roles_may_import(self):
        data = xlsx([["7502000000701", "X", "1", "", "", "", "", "", "", "", ""]])
        with self.assertRaises(AccessError):
            self.new_import(data, user=self.cashier)
        with self.assertRaises(AccessError):
            self.new_import(data, user=self.cashier_quick)
        record = self.new_import(data, user=self.inventory)
        record.action_validate()
        self.assertEqual(record.rows_to_create, 1)
        # Another company's user does not even see it.
        self.assertFalse(self.env["vlux.catalog.import"].with_user(self.owner_b).search([("id", "=", record.id)]))

    def test_file_cannot_change_once_imported(self):
        record = self.run_import(self.new_import(xlsx([["7502000000801", "Y", "1", "", "", "", "", "", "", "", ""]])))
        with self.assertRaises(UserError):
            record.write({"file": base64.b64encode(b"otro")})


@tagged("post_install", "-at_install")
class TestCatalogImportTemplate(HttpCase, VluxCatalogCase):

    def test_template_lists_the_company_taxes_and_is_restricted(self):
        import openpyxl

        self.owner.sudo().password = "owner-pass-123"
        self.authenticate("catalog-owner", "owner-pass-123")
        response = self.url_open("/vlux_pos_catalog/import/template.xlsx")
        self.assertEqual(response.status_code, 200)
        book = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.sheetnames, ["Productos", "Impuestos", "Categorias", "Instrucciones"])
        self.assertEqual([cell.value for cell in book["Productos"][1]], HEADERS)

        self.cashier.sudo().password = "cashier-pass-123"
        self.authenticate("catalog-cashier", "cashier-pass-123")
        response = self.url_open("/vlux_pos_catalog/import/template.xlsx")
        self.assertNotEqual(response.status_code, 200)
